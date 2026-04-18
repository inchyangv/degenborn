#!/usr/bin/env python3
"""
DegenBorn — Hackathon Pitch Deck Generator
Produces a branded 12-slide PPTX with cyberpunk neon aesthetic.

Brand palette:
  bg:       #0A0A0F
  card:     #111118
  green:    #00FF88
  purple:   #9945FF
  cyan:     #00D4FF
  gold:     #FFD700
  red:      #FF3D3D
  text:     #D0D0E8
  muted:    #6B6B8D
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ─── Palette ──────────────────────────────────────────────────────────
BG        = RGBColor(0x0A, 0x0A, 0x0F)
CARD      = RGBColor(0x11, 0x11, 0x18)
GREEN     = RGBColor(0x00, 0xFF, 0x88)
PURPLE    = RGBColor(0x99, 0x45, 0xFF)
CYAN      = RGBColor(0x00, 0xD4, 0xFF)
GOLD      = RGBColor(0xFF, 0xD7, 0x00)
RED       = RGBColor(0xFF, 0x3D, 0x3D)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
TEXT      = RGBColor(0xD0, 0xD0, 0xE8)
MUTED     = RGBColor(0x6B, 0x6B, 0x8D)
DIM       = RGBColor(0x2A, 0x2A, 0x3A)
BORDER    = RGBColor(0x1E, 0x1E, 0x2E)

# Archetype colors
ARCH_COLORS = {
    "Mad Gambler":     RGBColor(0xFF, 0x3D, 0x3D),
    "Ice Whale":       RGBColor(0x00, 0xD4, 0xFF),
    "Rug Necromancer": RGBColor(0x99, 0x45, 0xFF),
    "Diamond Cultist": RGBColor(0x88, 0xCC, 0xFF),
    "Sniper Jester":   RGBColor(0xFF, 0xD7, 0x00),
    "Ghost Bagholder": RGBColor(0xAA, 0xAA, 0xAA),
}

# ─── Fonts ────────────────────────────────────────────────────────────
FONT_TITLE = "Arial Black"
FONT_BODY  = "Calibri"
FONT_MONO  = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO_PATH = os.path.join(REPO, "LOGO.png")
ICON_PATH = os.path.join(REPO, "LOGO_ICON.png")
OUT_PATH  = os.path.join(REPO, "docs", "DegenBorn_Pitch.pptx")


def set_bg(slide, color=BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_accent_bar(slide, top=True, color=GREEN, thickness=Inches(0.04)):
    """Thin neon accent bar across the slide."""
    y = 0 if top else SLIDE_H - thickness
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0, y, SLIDE_W, thickness
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def add_glow_blob(slide, left, top, width, height, color, alpha=0.18):
    """Semi-transparent oval for ambient glow."""
    blob = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        left, top, width, height
    )
    blob.fill.solid()
    blob.fill.fore_color.rgb = color
    blob.line.fill.background()
    # python-pptx doesn't support alpha on shape fills directly,
    # so we approximate with a muted version of the color
    r = int(color[0] * alpha + BG[0] * (1 - alpha))
    g = int(color[1] * alpha + BG[1] * (1 - alpha))
    b = int(color[2] * alpha + BG[2] * (1 - alpha))
    blob.fill.fore_color.rgb = RGBColor(min(r, 255), min(g, 255), min(b, 255))


def add_corner_marks(slide, color=GREEN, length=Inches(0.5), thickness=Pt(2)):
    """L-shaped corner brackets — cyberpunk HUD style."""
    corners = [
        # top-left
        (Inches(0.3), Inches(0.3), Inches(0.3) + length, Inches(0.3)),
        (Inches(0.3), Inches(0.3), Inches(0.3), Inches(0.3) + length),
        # top-right
        (SLIDE_W - Inches(0.3) - length, Inches(0.3), SLIDE_W - Inches(0.3), Inches(0.3)),
        (SLIDE_W - Inches(0.3), Inches(0.3), SLIDE_W - Inches(0.3), Inches(0.3) + length),
        # bottom-left
        (Inches(0.3), SLIDE_H - Inches(0.3), Inches(0.3) + length, SLIDE_H - Inches(0.3)),
        (Inches(0.3), SLIDE_H - Inches(0.3) - length, Inches(0.3), SLIDE_H - Inches(0.3)),
        # bottom-right
        (SLIDE_W - Inches(0.3) - length, SLIDE_H - Inches(0.3), SLIDE_W - Inches(0.3), SLIDE_H - Inches(0.3)),
        (SLIDE_W - Inches(0.3), SLIDE_H - Inches(0.3) - length, SLIDE_W - Inches(0.3), SLIDE_H - Inches(0.3)),
    ]
    for x1, y1, x2, y2 in corners:
        line = slide.shapes.add_connector(1, x1, y1, x2, y2)  # MSO_CONNECTOR_TYPE.STRAIGHT
        line.line.color.rgb = color
        line.line.width = thickness


def add_text_box(slide, left, top, width, height, text, font_name=FONT_BODY,
                 font_size=Pt(16), color=TEXT, bold=False, italic=False,
                 alignment=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.3):
    txbox = slide.shapes.add_textbox(left, top, width, height)
    tf = txbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = font_name
    p.font.size = font_size
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.italic = italic
    p.alignment = alignment
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    if line_spacing:
        p.line_spacing = line_spacing
    tf.auto_size = None
    return txbox


def add_multiline_box(slide, left, top, width, height, lines, defaults=None):
    """
    lines: list of dicts with keys:
      text, font_name, font_size, color, bold, italic, alignment, space_after
    """
    txbox = slide.shapes.add_textbox(left, top, width, height)
    tf = txbox.text_frame
    tf.word_wrap = True
    tf.auto_size = None

    d = defaults or {}
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line.get("text", "")
        p.font.name = line.get("font_name", d.get("font_name", FONT_BODY))
        p.font.size = line.get("font_size", d.get("font_size", Pt(16)))
        p.font.color.rgb = line.get("color", d.get("color", TEXT))
        p.font.bold = line.get("bold", d.get("bold", False))
        p.font.italic = line.get("italic", False)
        p.alignment = line.get("alignment", d.get("alignment", PP_ALIGN.LEFT))
        p.space_before = Pt(0)
        p.space_after = line.get("space_after", Pt(4))
        p.line_spacing = line.get("line_spacing", 1.2)
    return txbox


def add_card(slide, left, top, width, height, color=CARD, border_color=BORDER):
    """Dark card panel with subtle border."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    # Adjust corner radius
    card.adjustments[0] = 0.04
    return card


def add_watermark(slide):
    """Small icon watermark in bottom-right."""
    if os.path.exists(ICON_PATH):
        slide.shapes.add_picture(
            ICON_PATH,
            SLIDE_W - Inches(1.0),
            SLIDE_H - Inches(1.0),
            Inches(0.6),
            Inches(0.6),
        )


def add_slide_number(slide, num, total=12):
    add_text_box(
        slide, SLIDE_W - Inches(1.2), SLIDE_H - Inches(0.5),
        Inches(1.0), Inches(0.3),
        f"{num:02d} / {total:02d}",
        font_name=FONT_MONO, font_size=Pt(9), color=MUTED,
        alignment=PP_ALIGN.RIGHT
    )


# ═════════════════════════════════════════════════════════════════════
# SLIDES
# ═════════════════════════════════════════════════════════════════════

def slide_01_cover(prs):
    """Title slide — logo, name, tagline."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_bg(slide)

    # Ambient glow blobs
    add_glow_blob(slide, Inches(-1), Inches(-1), Inches(6), Inches(6), GREEN)
    add_glow_blob(slide, Inches(9), Inches(3), Inches(5), Inches(5), PURPLE)

    add_accent_bar(slide, top=True, color=GREEN)
    add_accent_bar(slide, top=False, color=PURPLE)
    add_corner_marks(slide, GREEN)

    # Logo
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(
            LOGO_PATH,
            Inches(3.8), Inches(1.2),
            Inches(5.7), Inches(2.4),
        )

    # Tagline
    add_text_box(
        slide, Inches(1.5), Inches(4.2), Inches(10.3), Inches(1.0),
        "Your wallet has a personality. You just haven't seen it yet.",
        font_name=FONT_BODY, font_size=Pt(24), color=TEXT,
        bold=False, italic=True, alignment=PP_ALIGN.CENTER,
    )

    # Sub-info
    add_text_box(
        slide, Inches(1.5), Inches(5.5), Inches(10.3), Inches(0.5),
        "AI Identity Engine for Four.meme Traders  //  BNB Chain",
        font_name=FONT_MONO, font_size=Pt(13), color=MUTED,
        alignment=PP_ALIGN.CENTER,
    )

    # Hackathon badge
    add_text_box(
        slide, Inches(1.5), Inches(6.2), Inches(10.3), Inches(0.4),
        "Four.meme AI Sprint  |  April 2026  |  Solo Build",
        font_name=FONT_MONO, font_size=Pt(11), color=DIM,
        alignment=PP_ALIGN.CENTER,
    )


def slide_02_problem(prs):
    """The Problem — 3 pain points."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_accent_bar(slide, color=RED)
    add_glow_blob(slide, Inches(8), Inches(-2), Inches(6), Inches(6), RED)

    # Section label
    add_text_box(
        slide, Inches(0.8), Inches(0.5), Inches(4), Inches(0.4),
        "THE PROBLEM",
        font_name=FONT_MONO, font_size=Pt(12), color=RED, bold=True,
    )

    # Main headline
    add_text_box(
        slide, Inches(0.8), Inches(1.0), Inches(11), Inches(1.2),
        "On-chain data is rich with story.\nBut it collapses into a log nobody shares.",
        font_name=FONT_TITLE, font_size=Pt(32), color=WHITE, bold=True,
        alignment=PP_ALIGN.LEFT,
    )

    # Three problem cards
    problems = [
        ("01", "Wallets have no identity",
         "Your wallet shows what happened,\nnot who you are. No style, no scars,\nno personality."),
        ("02", "On-chain data is boring",
         "Transaction hashes and P&L tables.\nNobody screenshots a block explorer.\nZero shareability."),
        ("03", "NFTs are static",
         "Minted once, forgotten forever.\nNo growth. No connection to\nyour actual behavior."),
    ]

    for i, (num, title, desc) in enumerate(problems):
        x = Inches(0.8) + Inches(4.0) * i
        y = Inches(3.0)

        add_card(slide, x, y, Inches(3.6), Inches(3.5))

        # Number
        add_text_box(
            slide, x + Inches(0.3), y + Inches(0.3), Inches(1), Inches(0.5),
            num, font_name=FONT_TITLE, font_size=Pt(36), color=RED, bold=True,
        )
        # Title
        add_text_box(
            slide, x + Inches(0.3), y + Inches(1.0), Inches(3.0), Inches(0.6),
            title, font_name=FONT_BODY, font_size=Pt(17), color=WHITE, bold=True,
        )
        # Desc
        add_text_box(
            slide, x + Inches(0.3), y + Inches(1.7), Inches(3.0), Inches(1.6),
            desc, font_name=FONT_BODY, font_size=Pt(13), color=MUTED,
            line_spacing=1.4,
        )

    add_watermark(slide)
    add_slide_number(slide, 2)


def slide_03_insight(prs):
    """The Insight — big quote."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)

    add_glow_blob(slide, Inches(3), Inches(1), Inches(7), Inches(5), GREEN)
    add_corner_marks(slide, GREEN)

    # Big quote
    add_text_box(
        slide, Inches(1.5), Inches(2.0), Inches(10.3), Inches(2.0),
        "We don't need another dashboard.\nWe need a mirror.",
        font_name=FONT_TITLE, font_size=Pt(44), color=GREEN, bold=True,
        alignment=PP_ALIGN.CENTER,
    )

    # Explanation
    add_text_box(
        slide, Inches(2.5), Inches(4.5), Inches(8.3), Inches(1.5),
        "Panic sells at 3am. Diamond hands through -90%. Aping back in\n"
        "after getting rugged. That's not just data. That's personality.\n"
        "DegenBorn makes it visible.",
        font_name=FONT_BODY, font_size=Pt(18), color=TEXT,
        alignment=PP_ALIGN.CENTER, line_spacing=1.5,
    )

    add_slide_number(slide, 3)


def slide_04_how_it_works(prs):
    """Pipeline overview."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_accent_bar(slide, color=CYAN)

    add_text_box(
        slide, Inches(0.8), Inches(0.5), Inches(4), Inches(0.4),
        "HOW IT WORKS",
        font_name=FONT_MONO, font_size=Pt(12), color=CYAN, bold=True,
    )

    add_text_box(
        slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.8),
        "From wallet activity to living character",
        font_name=FONT_TITLE, font_size=Pt(34), color=WHITE, bold=True,
    )

    # Pipeline steps
    steps = [
        ("CONNECT", "Read Four.meme\nwallet activity", GREEN),
        ("ANALYZE", "Compute 5-axis\nPersona DNA", CYAN),
        ("CLASSIFY", "Rule-tree assigns\narchetype", PURPLE),
        ("GENERATE", "AI creates genesis\ncharacter", GOLD),
        ("EVOLVE", "Every trade mutates\nyour monster", GREEN),
    ]

    for i, (label, desc, color) in enumerate(steps):
        x = Inches(0.6) + Inches(2.5) * i
        y = Inches(2.5)

        # Step card
        add_card(slide, x, y, Inches(2.2), Inches(3.2))

        # Step number circle
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, x + Inches(0.75), y + Inches(0.3),
            Inches(0.7), Inches(0.7)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()

        add_text_box(
            slide, x + Inches(0.75), y + Inches(0.35),
            Inches(0.7), Inches(0.6),
            str(i + 1),
            font_name=FONT_TITLE, font_size=Pt(22), color=BG, bold=True,
            alignment=PP_ALIGN.CENTER,
        )

        # Label
        add_text_box(
            slide, x + Inches(0.2), y + Inches(1.2), Inches(1.8), Inches(0.5),
            label, font_name=FONT_MONO, font_size=Pt(13), color=color, bold=True,
            alignment=PP_ALIGN.CENTER,
        )

        # Description
        add_text_box(
            slide, x + Inches(0.2), y + Inches(1.8), Inches(1.8), Inches(1.2),
            desc, font_name=FONT_BODY, font_size=Pt(12), color=MUTED,
            alignment=PP_ALIGN.CENTER, line_spacing=1.4,
        )

        # Arrow between steps
        if i < len(steps) - 1:
            arrow_x = x + Inches(2.2)
            add_text_box(
                slide, arrow_x, y + Inches(1.2),
                Inches(0.3), Inches(0.5),
                ">",
                font_name=FONT_MONO, font_size=Pt(20), color=DIM,
                alignment=PP_ALIGN.CENTER,
            )

    # Bottom note
    add_text_box(
        slide, Inches(0.8), Inches(6.2), Inches(11), Inches(0.5),
        "Same wallet + same events = same result. Always. No LLM hallucination in the scoring layer.",
        font_name=FONT_MONO, font_size=Pt(12), color=GREEN, italic=True,
        alignment=PP_ALIGN.CENTER,
    )

    add_watermark(slide)
    add_slide_number(slide, 4)


def slide_05_persona_dna(prs):
    """Persona DNA — 5 axes."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_accent_bar(slide, color=GREEN)
    add_glow_blob(slide, Inches(9), Inches(0), Inches(5), Inches(5), GREEN)

    add_text_box(
        slide, Inches(0.8), Inches(0.5), Inches(4), Inches(0.4),
        "PERSONA DNA",
        font_name=FONT_MONO, font_size=Pt(12), color=GREEN, bold=True,
    )

    add_text_box(
        slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.8),
        "Five axes. Pure math. Zero LLM calls.",
        font_name=FONT_TITLE, font_size=Pt(34), color=WHITE, bold=True,
    )

    axes = [
        ("AGGRESSION", "82", "How fast and often you trade",    RED),
        ("CONVICTION", "34", "How long and focused you hold",   CYAN),
        ("CHAOS",      "76", "How volatile your history is",    PURPLE),
        ("LUCK",       "41", "How well-timed your exits are",   GOLD),
        ("SURVIVAL",   "91", "How many times you came back",    GREEN),
    ]

    for i, (name, score, desc, color) in enumerate(axes):
        y = Inches(2.2) + Inches(1.0) * i

        # Label
        add_text_box(
            slide, Inches(0.8), y, Inches(2.2), Inches(0.4),
            name, font_name=FONT_MONO, font_size=Pt(15), color=color, bold=True,
        )

        # Bar background
        bar_bg = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(3.2), y + Inches(0.05), Inches(6.0), Inches(0.35)
        )
        bar_bg.fill.solid()
        bar_bg.fill.fore_color.rgb = DIM
        bar_bg.line.fill.background()
        bar_bg.adjustments[0] = 0.5

        # Bar fill
        fill_w = Inches(6.0) * int(score) / 100
        if fill_w > 0:
            bar_fill = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(3.2), y + Inches(0.05), int(fill_w), Inches(0.35)
            )
            bar_fill.fill.solid()
            bar_fill.fill.fore_color.rgb = color
            bar_fill.line.fill.background()
            bar_fill.adjustments[0] = 0.5

        # Score
        add_text_box(
            slide, Inches(9.5), y, Inches(0.8), Inches(0.4),
            score, font_name=FONT_TITLE, font_size=Pt(20), color=color, bold=True,
            alignment=PP_ALIGN.CENTER,
        )

        # Description
        add_text_box(
            slide, Inches(10.5), y, Inches(2.5), Inches(0.4),
            desc, font_name=FONT_BODY, font_size=Pt(11), color=MUTED,
        )

    # Example label
    add_text_box(
        slide, Inches(0.8), Inches(7.0), Inches(5), Inches(0.3),
        'Example: "Rug Necromancer" profile',
        font_name=FONT_MONO, font_size=Pt(10), color=MUTED, italic=True,
    )

    add_slide_number(slide, 5)


def slide_06_archetypes(prs):
    """6 Archetypes grid."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_accent_bar(slide, color=PURPLE)

    add_text_box(
        slide, Inches(0.8), Inches(0.5), Inches(4), Inches(0.4),
        "ARCHETYPES",
        font_name=FONT_MONO, font_size=Pt(12), color=PURPLE, bold=True,
    )

    add_text_box(
        slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.8),
        "A rule tree decides. Not a prompt.",
        font_name=FONT_TITLE, font_size=Pt(34), color=WHITE, bold=True,
    )

    archetypes = [
        ("Mad Gambler",     "AGG >= 65, CHAOS >= 65",  "All-in, always. The casino\nnever closes if you never leave."),
        ("Ice Whale",       "CONV >= 65, LUCK >= 65",  "Patient, loaded, lucky.\nYou're not gambling. You're investing\nin chaos."),
        ("Rug Necromancer", "CHAOS >= 65, SURV >= 65",  "Your portfolio died.\nYou didn't. Death is just a dip."),
        ("Diamond Cultist", "CONV >= 65, LUCK < 35",   "Bought the top. Still holding.\nNot because it'll moon. Because\nleaving means admitting defeat."),
        ("Sniper Jester",   "AGG >= 65, LUCK >= 65",   "In and out. Perfectly timed.\nNobody knows if you're a genius\nor just lucky."),
        ("Ghost Bagholder", "CONV >= 65, SURV < 35",   "You believed. Hard.\nThe token didn't. Still haunting\nthe chart."),
    ]

    for i, (name, rule, desc) in enumerate(archetypes):
        col = i % 3
        row = i // 3
        x = Inches(0.6) + Inches(4.1) * col
        y = Inches(2.2) + Inches(2.6) * row

        color = ARCH_COLORS[name]
        add_card(slide, x, y, Inches(3.8), Inches(2.3))

        # Color dot
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, x + Inches(0.25), y + Inches(0.3),
            Inches(0.2), Inches(0.2)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()

        # Name
        add_text_box(
            slide, x + Inches(0.55), y + Inches(0.25),
            Inches(3.0), Inches(0.35),
            name, font_name=FONT_BODY, font_size=Pt(16), color=color, bold=True,
        )

        # Rule
        add_text_box(
            slide, x + Inches(0.25), y + Inches(0.7),
            Inches(3.3), Inches(0.3),
            rule, font_name=FONT_MONO, font_size=Pt(9), color=MUTED,
        )

        # Description
        add_text_box(
            slide, x + Inches(0.25), y + Inches(1.1),
            Inches(3.3), Inches(1.1),
            desc, font_name=FONT_BODY, font_size=Pt(11), color=TEXT,
            line_spacing=1.3,
        )

    add_watermark(slide)
    add_slide_number(slide, 6)


def slide_07_evolution(prs):
    """Evolution — how characters change."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_accent_bar(slide, color=GOLD)
    add_glow_blob(slide, Inches(-1), Inches(3), Inches(5), Inches(5), GOLD)

    add_text_box(
        slide, Inches(0.8), Inches(0.5), Inches(4), Inches(0.4),
        "EVOLUTION",
        font_name=FONT_MONO, font_size=Pt(12), color=GOLD, bold=True,
    )

    add_text_box(
        slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.8),
        "Your monster isn't static. Every trade leaves a mark.",
        font_name=FONT_TITLE, font_size=Pt(32), color=WHITE, bold=True,
    )

    # Event -> Mutation pairs
    events = [
        ("Win streak x3",        "Crown appears",           GREEN,  "+1 crownCount, mood: euphoria"),
        ("Major loss",           "Scar added",              RED,    "+1 scarCount, mood: despair"),
        ("Rug pull detected",    "Zombie eyes activate",    PURPLE, "+20 corruption, zombieTrait: true"),
        ("Comeback after -90%",  "Revenge aura ignites",    GOLD,   "+1 survivalStreak, mood: revenge"),
        ("Long diamond hold",    "Royal cloak forms",       CYAN,   "+15 prestige, royalTrait: true"),
        ("Multiple rug survival","Necromancer ascension",   PURPLE, "Full re-render triggered"),
    ]

    # Left column: Event
    add_text_box(
        slide, Inches(1.0), Inches(2.2), Inches(4), Inches(0.4),
        "EVENT", font_name=FONT_MONO, font_size=Pt(11), color=MUTED, bold=True,
    )
    # Right column: Mutation
    add_text_box(
        slide, Inches(5.5), Inches(2.2), Inches(4), Inches(0.4),
        "MUTATION", font_name=FONT_MONO, font_size=Pt(11), color=MUTED, bold=True,
    )
    # Far right: State change
    add_text_box(
        slide, Inches(9.5), Inches(2.2), Inches(3.5), Inches(0.4),
        "STATE DELTA", font_name=FONT_MONO, font_size=Pt(11), color=MUTED, bold=True,
    )

    for i, (event, mutation, color, delta) in enumerate(events):
        y = Inches(2.8) + Inches(0.7) * i

        # Thin separator
        sep = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.8), y - Inches(0.05), Inches(12.0), Pt(1)
        )
        sep.fill.solid()
        sep.fill.fore_color.rgb = DIM
        sep.line.fill.background()

        add_text_box(
            slide, Inches(1.0), y, Inches(4.0), Inches(0.5),
            event, font_name=FONT_BODY, font_size=Pt(15), color=WHITE, bold=True,
        )

        # Arrow
        add_text_box(
            slide, Inches(5.0), y, Inches(0.4), Inches(0.5),
            ">", font_name=FONT_MONO, font_size=Pt(15), color=color, bold=True,
        )

        add_text_box(
            slide, Inches(5.5), y, Inches(3.8), Inches(0.5),
            mutation, font_name=FONT_BODY, font_size=Pt(15), color=color, bold=True,
        )

        add_text_box(
            slide, Inches(9.5), y, Inches(3.5), Inches(0.5),
            delta, font_name=FONT_MONO, font_size=Pt(10), color=MUTED,
        )

    # Bottom note
    add_text_box(
        slide, Inches(0.8), Inches(6.8), Inches(12), Inches(0.4),
        "Genesis image generated once. Daily changes use overlay compositing. Full re-render only at major milestones.",
        font_name=FONT_MONO, font_size=Pt(11), color=GREEN, italic=True,
    )

    add_slide_number(slide, 7)


def slide_08_soul_core(prs):
    """Soul Core — Soulbound NFT."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_accent_bar(slide, color=CYAN)
    add_glow_blob(slide, Inches(6), Inches(1), Inches(6), Inches(6), CYAN)

    add_text_box(
        slide, Inches(0.8), Inches(0.5), Inches(4), Inches(0.4),
        "SOUL CORE",
        font_name=FONT_MONO, font_size=Pt(12), color=CYAN, bold=True,
    )

    add_text_box(
        slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.8),
        "Your on-chain identity. Not a collectible.",
        font_name=FONT_TITLE, font_size=Pt(34), color=WHITE, bold=True,
    )

    # Two-column layout
    # Left: Soul Core
    add_card(slide, Inches(0.8), Inches(2.3), Inches(5.5), Inches(4.5))

    add_text_box(
        slide, Inches(1.2), Inches(2.6), Inches(4.7), Inches(0.5),
        "SOUL CORE", font_name=FONT_MONO, font_size=Pt(20), color=CYAN, bold=True,
    )

    soul_core_lines = [
        "Soulbound ERC-721 on BNB Chain",
        "One per wallet. Non-transferable.",
        "Always reflects your current state.",
        "",
        "On-chain data:",
        "  dnaHash      // persona fingerprint",
        "  archetype    // rule-tree result",
        "  level        // growth stage",
        "  stateHash    // latest mutation",
    ]
    for i, line in enumerate(soul_core_lines):
        y = Inches(3.3) + Inches(0.32) * i
        color = MUTED if line.startswith("  ") else TEXT
        if line.startswith("On-chain"):
            color = CYAN
        add_text_box(
            slide, Inches(1.2), y, Inches(4.7), Inches(0.3),
            line,
            font_name=FONT_MONO if line.startswith("  ") or line.startswith("On-chain") else FONT_BODY,
            font_size=Pt(12) if line.startswith("  ") else Pt(14),
            color=color,
        )

    # Right: Snapshot Relic
    add_card(slide, Inches(7.0), Inches(2.3), Inches(5.5), Inches(4.5))

    add_text_box(
        slide, Inches(7.4), Inches(2.6), Inches(4.7), Inches(0.5),
        "SNAPSHOT RELIC", font_name=FONT_MONO, font_size=Pt(20), color=GOLD, bold=True,
    )

    relic_items = [
        ("First Crowned Win",      "Win streak milestone"),
        ("Rug Survivor",           "Survived 3+ rug events"),
        ("Seven-Day Resurrection", "Major comeback"),
        ("Chaos Ascension",        "Corruption threshold reached"),
    ]
    for i, (title, desc) in enumerate(relic_items):
        y = Inches(3.4) + Inches(0.7) * i
        # Dot
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(7.4), y + Inches(0.05),
            Inches(0.15), Inches(0.15)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = GOLD
        dot.line.fill.background()

        add_text_box(
            slide, Inches(7.7), y, Inches(4.5), Inches(0.35),
            title, font_name=FONT_BODY, font_size=Pt(15), color=WHITE, bold=True,
        )
        add_text_box(
            slide, Inches(7.7), y + Inches(0.35), Inches(4.5), Inches(0.3),
            desc, font_name=FONT_BODY, font_size=Pt(11), color=MUTED,
        )

    # Bottom note
    add_text_box(
        slide, Inches(7.4), Inches(6.2), Inches(5.0), Inches(0.4),
        "Tradeable. Each one commemorates a major evolution moment.",
        font_name=FONT_BODY, font_size=Pt(12), color=MUTED, italic=True,
    )

    add_watermark(slide)
    add_slide_number(slide, 8)


def slide_09_edge(prs):
    """Why This Has Edge — 3 differentiators."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)

    add_glow_blob(slide, Inches(0), Inches(0), Inches(5), Inches(5), GREEN)
    add_glow_blob(slide, Inches(8), Inches(3), Inches(5), Inches(5), PURPLE)
    add_corner_marks(slide, GREEN)

    # Big statement
    add_text_box(
        slide, Inches(0.8), Inches(0.8), Inches(11.5), Inches(1.2),
        "Rules decide.\nAI expresses.",
        font_name=FONT_TITLE, font_size=Pt(48), color=GREEN, bold=True,
    )

    # Three differentiators
    diffs = [
        (
            "Deterministic Identity",
            "Most AI x Web3 projects let the LLM make judgments.\nWe don't. Scoring engine and archetype classifier are\nfully deterministic. AI only handles narrative, visuals,\nand captions. Reproducible. Auditable. Stable.",
            GREEN,
        ),
        (
            "Trading Becomes Retention",
            "Every trade changes your character. That's not a gimmick\n--- it's a feedback loop. Four.meme gets stickier\nwhen your portfolio has a face, scars, and a mood.",
            CYAN,
        ),
        (
            "Sharing Without Spam",
            "We generate meme-ready share cards and one-line\ncaptions. Users decide when to post. No auto-tweeting.\nNo bot behavior. The content markets itself.",
            GOLD,
        ),
    ]

    for i, (title, desc, color) in enumerate(diffs):
        x = Inches(0.6) + Inches(4.1) * i
        y = Inches(3.0)

        add_card(slide, x, y, Inches(3.8), Inches(3.8))

        # Color accent bar on card
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            x, y, Inches(0.06), Inches(3.8)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        add_text_box(
            slide, x + Inches(0.3), y + Inches(0.3), Inches(3.2), Inches(0.5),
            title, font_name=FONT_BODY, font_size=Pt(18), color=color, bold=True,
        )

        add_text_box(
            slide, x + Inches(0.3), y + Inches(1.0), Inches(3.2), Inches(2.5),
            desc, font_name=FONT_BODY, font_size=Pt(12), color=TEXT,
            line_spacing=1.5,
        )

    add_slide_number(slide, 9)


def slide_10_architecture(prs):
    """Tech Stack."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_accent_bar(slide, color=CYAN)

    add_text_box(
        slide, Inches(0.8), Inches(0.5), Inches(4), Inches(0.4),
        "ARCHITECTURE",
        font_name=FONT_MONO, font_size=Pt(12), color=CYAN, bold=True,
    )

    add_text_box(
        slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.8),
        "Built to ship, not to impress a whitepaper.",
        font_name=FONT_TITLE, font_size=Pt(32), color=WHITE, bold=True,
    )

    # Tech stack as a table-like layout
    stack = [
        ("Frontend",       "Next.js 14, TypeScript, Tailwind, wagmi v2",     GREEN),
        ("Scoring",        "Deterministic engine — pure functions, 0 LLM",   GREEN),
        ("Classification", "Rule-tree archetype mapper",                     CYAN),
        ("Narrative",      "Claude / GPT-4o-mini + structured output",       PURPLE),
        ("Image",          "DALL-E 3 genesis + Canvas overlay compositor",   PURPLE),
        ("Contract",       "Solidity 0.8.24, Soulbound ERC-721, BNB Chain", GOLD),
        ("Data",           "Moralis / Covalent + offline fixture mode",      CYAN),
        ("Persistence",    "Postgres + Vercel Blob",                         CYAN),
    ]

    for i, (layer, tech, color) in enumerate(stack):
        y = Inches(2.3) + Inches(0.6) * i

        # Separator
        if i > 0:
            sep = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(0.8), y - Inches(0.05), Inches(11.7), Pt(1)
            )
            sep.fill.solid()
            sep.fill.fore_color.rgb = DIM
            sep.line.fill.background()

        # Color dot
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.9), y + Inches(0.08),
            Inches(0.15), Inches(0.15)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()

        add_text_box(
            slide, Inches(1.2), y, Inches(2.5), Inches(0.45),
            layer, font_name=FONT_MONO, font_size=Pt(14), color=color, bold=True,
        )

        add_text_box(
            slide, Inches(4.0), y, Inches(8.5), Inches(0.45),
            tech, font_name=FONT_BODY, font_size=Pt(14), color=TEXT,
        )

    # Monorepo note
    add_text_box(
        slide, Inches(0.8), Inches(6.8), Inches(12), Inches(0.4),
        "pnpm monorepo  //  apps/web + apps/worker + packages/shared + packages/contracts",
        font_name=FONT_MONO, font_size=Pt(11), color=MUTED, italic=True,
    )

    add_watermark(slide)
    add_slide_number(slide, 10)


def slide_11_whats_live(prs):
    """What's Live — feature checklist."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_accent_bar(slide, color=GREEN)
    add_glow_blob(slide, Inches(9), Inches(0), Inches(5), Inches(5), GREEN)

    add_text_box(
        slide, Inches(0.8), Inches(0.5), Inches(4), Inches(0.4),
        "WHAT'S LIVE",
        font_name=FONT_MONO, font_size=Pt(12), color=GREEN, bold=True,
    )

    add_text_box(
        slide, Inches(0.8), Inches(1.0), Inches(11), Inches(0.8),
        "This isn't a concept. It's deployed.",
        font_name=FONT_TITLE, font_size=Pt(34), color=WHITE, bold=True,
    )

    features_left = [
        "Wallet connect  >  DNA analysis in <30s",
        "AI genesis character generation",
        "12 overlay traits (crown, scar, zombie eyes...)",
        "9-event state machine",
        "Mutation Diary tracking every change",
        "Meme-ready share cards with captions",
    ]

    features_right = [
        "Soulbound Soul Core on BNB Chain",
        "Replay Mode — full offline demo",
        "30+ pages: Gallery, Zodiac, Graveyard...",
        "Report Card with detailed stats",
        "Hall of Fame + Battle comparisons",
        "Token launch via Four.meme integration",
    ]

    for i, feat in enumerate(features_left):
        y = Inches(2.3) + Inches(0.65) * i

        # Checkmark
        add_text_box(
            slide, Inches(0.8), y, Inches(0.5), Inches(0.4),
            "+", font_name=FONT_MONO, font_size=Pt(16), color=GREEN, bold=True,
        )
        add_text_box(
            slide, Inches(1.3), y, Inches(5.0), Inches(0.5),
            feat, font_name=FONT_BODY, font_size=Pt(15), color=TEXT,
        )

    for i, feat in enumerate(features_right):
        y = Inches(2.3) + Inches(0.65) * i

        add_text_box(
            slide, Inches(7.0), y, Inches(0.5), Inches(0.4),
            "+", font_name=FONT_MONO, font_size=Pt(16), color=GREEN, bold=True,
        )
        add_text_box(
            slide, Inches(7.5), y, Inches(5.0), Inches(0.5),
            feat, font_name=FONT_BODY, font_size=Pt(15), color=TEXT,
        )

    # Bottom emphasis
    add_card(slide, Inches(2.5), Inches(6.3), Inches(8.3), Inches(0.7))
    add_text_box(
        slide, Inches(2.5), Inches(6.35), Inches(8.3), Inches(0.6),
        "Solo build.  12 days.  Fully working product.",
        font_name=FONT_TITLE, font_size=Pt(22), color=GREEN, bold=True,
        alignment=PP_ALIGN.CENTER,
    )

    add_slide_number(slide, 11)


def slide_12_close(prs):
    """Closing slide — CTA."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)

    add_glow_blob(slide, Inches(2), Inches(0), Inches(9), Inches(7), GREEN)
    add_glow_blob(slide, Inches(7), Inches(2), Inches(6), Inches(6), PURPLE)
    add_accent_bar(slide, top=True, color=GREEN)
    add_accent_bar(slide, top=False, color=PURPLE)
    add_corner_marks(slide, GREEN)

    # Logo
    if os.path.exists(ICON_PATH):
        slide.shapes.add_picture(
            ICON_PATH,
            Inches(5.9), Inches(0.8),
            Inches(1.5), Inches(1.5),
        )

    # Main statement
    add_text_box(
        slide, Inches(1.5), Inches(2.6), Inches(10.3), Inches(1.5),
        "DegenBorn is not another NFT project.",
        font_name=FONT_TITLE, font_size=Pt(38), color=WHITE, bold=True,
        alignment=PP_ALIGN.CENTER,
    )

    add_text_box(
        slide, Inches(1.5), Inches(3.8), Inches(10.3), Inches(1.2),
        "It's the first time your wallet behavior\nbecomes a character with scars, crowns, and a mood.",
        font_name=FONT_BODY, font_size=Pt(22), color=TEXT,
        alignment=PP_ALIGN.CENTER, line_spacing=1.5,
    )

    # Value prop
    add_text_box(
        slide, Inches(1.5), Inches(5.2), Inches(10.3), Inches(0.6),
        "Four.meme gets retention.  The user gets a self.",
        font_name=FONT_BODY, font_size=Pt(20), color=GREEN,
        alignment=PP_ALIGN.CENTER, bold=True,
    )

    # CTA
    add_text_box(
        slide, Inches(1.5), Inches(6.2), Inches(10.3), Inches(0.5),
        "degenborn.xyz",
        font_name=FONT_MONO, font_size=Pt(20), color=CYAN,
        alignment=PP_ALIGN.CENTER, bold=True,
    )

    add_text_box(
        slide, Inches(1.5), Inches(6.8), Inches(10.3), Inches(0.4),
        "Your wallet, reborn as a monster.",
        font_name=FONT_BODY, font_size=Pt(16), color=MUTED,
        alignment=PP_ALIGN.CENTER, italic=True,
    )


# ═════════════════════════════════════════════════════════════════════
# MAIN
# ═════════════════════════════════════════════════════════════════════

def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_01_cover(prs)
    slide_02_problem(prs)
    slide_03_insight(prs)
    slide_04_how_it_works(prs)
    slide_05_persona_dna(prs)
    slide_06_archetypes(prs)
    slide_07_evolution(prs)
    slide_08_soul_core(prs)
    slide_09_edge(prs)
    slide_10_architecture(prs)
    slide_11_whats_live(prs)
    slide_12_close(prs)

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    prs.save(OUT_PATH)
    print(f"Saved: {OUT_PATH}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
